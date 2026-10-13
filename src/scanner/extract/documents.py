"""DOCX, EML, HTML/XML, SQL dump and PPTX extraction into TextUnits (location {} for all of them).

Every extractor keeps the structure that gives values their meaning: table cells carry their column header as
`key`, form fields their label, and e-mail addresses come back as "From: Name <addr>" lines, so the detector sees
the same "key next to value" patterns it sees in logs and configs.
"""

import email
import re
import zipfile
from email import policy

from ..units import TextUnit
from .tables import table_units

# ------------------------------------------------------------------ DOCX


def docx_units(path):
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(path)
    units, lines, order = [], [], 0

    def flush():
        nonlocal lines, order
        if any(l.strip() for l in lines):
            units.append(TextUnit("\n".join(lines), {}, "docx-text", use_ner=True, order=order))
            order += 1
        lines = []

    # Walk the body in document order (python-docx's d.paragraphs and d.tables would separate them). Consecutive
    # paragraphs become one text unit; each table becomes header-keyed cells.
    for el in d.element.body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            lines.append(Paragraph(el, d).text)
        elif tag == "tbl":
            flush()
            t = Table(el, d)
            rows = [[c.text for c in r.cells] for r in t.rows]
            units.extend(table_units(rows, {}, "docx-table", order=order))
            order += 1
    flush()
    for s in d.sections:  # page headers and footers carry author names and contact lines too
        for part in (s.header, s.footer):
            try:
                txt = "\n".join(p.text for p in part.paragraphs)
            except Exception:
                txt = ""
            if txt.strip():
                units.append(TextUnit(txt, {}, "docx-header-footer", order=order))
                order += 1
    return units


# ------------------------------------------------------------------ PPTX (text only)


def pptx_units(path):
    # Slides are read straight from the XML (<a:p> paragraphs of <a:t> runs), so there is no python-pptx dependency.
    # The sort is lexicographic (slide10 before slide2); only the output order depends on it.
    units = []
    with zipfile.ZipFile(path) as z:
        slides = sorted(n for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n))
        for i, n in enumerate(slides):
            xml = z.read(n).decode("utf-8", "replace")
            paras = re.findall(r"<a:p>(.*?)</a:p>", xml, re.S)
            text = "\n".join("".join(re.findall(r"<a:t>([^<]*)</a:t>", p)) for p in paras)
            if text.strip():
                units.append(TextUnit(_unescape_xml(text), {}, "pptx-text", use_ner=True, order=i))
    return units


def _unescape_xml(s):
    import html

    return html.unescape(s)


# ------------------------------------------------------------------ EML

ADDRESS_HEADERS = ("From", "To", "Cc", "Bcc", "Reply-To", "Sender")


def eml_units(path=None, raw=None, depth=0):
    """Units for an e-mail: one for the address headers and the subject, then one per text part. Attached messages
    are recursed into (up to depth 3) and other attachments are sniffed and extracted like files."""
    if raw is None:
        with open(path, "rb") as f:
            raw = f.read()
    # policy.default decodes RFC 2047 encoded words in headers and gives structured address objects
    msg = email.message_from_bytes(raw, policy=policy.default)
    lines = []
    # Only address headers and the subject are kept: Message-ID, References and Received are machine identifiers
    # that look like e-mail addresses ("<3b9f...@mail.example>") but are not labelled as such.
    for h in ADDRESS_HEADERS:
        for v in msg.get_all(h, []) or []:
            try:
                addrs = v.addresses
            except AttributeError:
                addrs = ()
            if addrs:
                # one normalised line per address: quoting and encoded-word noise disappear, and the display name
                # gets the strongest name context there is ("From: Name <addr>")
                for a in addrs:
                    name = (a.display_name or "").strip()
                    lines.append(f"{h}: {name} <{a.addr_spec}>" if name else f"{h}: {a.addr_spec}")
            else:
                lines.append(f"{h}: {v}")
    if msg.get("Subject"):
        lines.append(f"Subject: {msg.get('Subject')}")
    units = [TextUnit("\n".join(lines), {}, "eml-headers", order=0)] if lines else []
    order = 1
    for part in msg.walk():
        if part.is_multipart():
            continue
        ctype = part.get_content_type()
        cte = (part.get("Content-Transfer-Encoding") or "7bit").lower()
        fname = part.get_filename()
        try:
            if ctype in ("text/plain", "text/html") and not fname:
                # get_content() undoes base64 / quoted-printable and applies the declared charset. QP's soft line
                # breaks ("11/=\n29") would otherwise split values in two. The fallback covers a wrong or
                # unknown charset.
                try:
                    body = part.get_content()
                except Exception:
                    body = (part.get_payload(decode=True) or b"").decode("utf-8", "replace")
                if ctype == "text/html":
                    for u in html_units(body):
                        u.extractor = f"eml-html:{cte}"
                        u.order = order
                        units.append(u)
                else:
                    units.append(TextUnit(body, {}, f"eml-body:{cte}", use_ner=True, order=order))
            elif ctype == "message/rfc822" and depth < 3:
                payload = part.get_payload()
                inner = payload[0].as_bytes() if isinstance(payload, list) and payload else part.get_payload(decode=True)
                if inner:
                    units.extend(eml_units(raw=inner, depth=depth + 1))
            else:
                data = part.get_payload(decode=True)
                if data:
                    units.extend(attachment_units(data, cte))
        except Exception:  # one undecodable part must not cost the rest of the message
            continue
        order += 1
    return units


def attachment_units(data, cte):
    """Text-bearing attachments are scanned as part of the e-mail (location {})."""
    import os
    import tempfile

    from ..filetype import decode_bytes, sniff

    # sniff() and the DOCX / PDF extractors work on paths, so the attachment is written to a temporary file. It is
    # closed before being reopened, because Windows does not allow a second open while the first is active.
    with tempfile.NamedTemporaryFile(delete=False) as tf:
        tf.write(data)
        tmp = tf.name
    try:
        s = sniff(tmp)
        if s.error:
            return []
        if s.kind == "docx":
            return docx_units(tmp)
        if s.kind == "pdf":
            from .pdf import pdf_text_only

            return [TextUnit(t, {}, "eml-attachment-pdf") for t in pdf_text_only(tmp)]
        if s.kind in ("text", "csv", "json", "jsonl", "sql", "xml"):
            return [TextUnit(decode_bytes(data)[0], {}, f"eml-attachment:{cte}")]
        if s.kind == "html":
            return html_units(decode_bytes(data)[0])
        return []
    finally:
        try:
            os.unlink(tmp)
        except OSError:  # Windows refuses to delete a file that a library still holds open; keep the findings
            pass


# ------------------------------------------------------------------ HTML / XML


def html_units(html_text, xml=False):
    import warnings

    from bs4 import BeautifulSoup, Comment

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        soup = BeautifulSoup(html_text, "lxml-xml" if xml else "lxml")
    units, order = [], 0
    for t in soup(["style", "noscript", "svg"]):  # presentation only, nothing to find
        t.decompose()
    # Inline scripts and comments are scanned as their own units rather than thrown away: they are a classic place
    # for leftover API keys and commented-out credentials. Each piece is removed from the tree once handled, so the
    # final get_text() never reports it a second time.
    for s in soup("script"):
        code = s.get_text()
        if code.strip():
            units.append(TextUnit(code, {}, "html-script", order=order))
            order += 1
        s.decompose()
    for c in soup.find_all(string=lambda t: isinstance(t, Comment)):
        if c.strip():
            units.append(TextUnit(str(c), {}, "html-comment", order=order))
            order += 1
        c.extract()
    if not xml:
        for table in soup.find_all("table"):
            # nested layout tables are handled as part of their outermost table
            if table.find_parent("table") is not None:
                continue
            rows = []
            for tr in table.find_all("tr"):
                rows.append([cell.get_text(" ", strip=True) for cell in tr.find_all(["th", "td"])])
            units.extend(table_units(rows, {}, "html-table", order=order))
            order += 1
            table.decompose()
        # A form field's key is its visible <label for=...> when there is one, otherwise its name or id attribute.
        labels = {lab.get("for"): lab.get_text(" ", strip=True) for lab in soup.find_all("label") if lab.get("for")}
        for inp in soup.find_all(["input", "textarea"]):
            val = inp.get("value") if inp.name == "input" else inp.get_text()
            # hidden inputs are scanned too: they are where pages tend to carry tokens and session values
            if val and val.strip() and (inp.get("type") or "text").lower() not in ("submit", "button", "checkbox",
                                                                                     "radio"):
                key = labels.get(inp.get("id")) or inp.get("name") or inp.get("id") or ""
                units.append(TextUnit(val.strip(), {}, "html-form", key=key, cell=True, order=order))
                order += 1
        for dl in soup.find_all("dl"):  # definition lists become "term: description" lines, i.e. key: value
            pairs = []
            for dt in dl.find_all("dt"):
                dd = dt.find_next_sibling("dd")
                if dd is not None:
                    pairs.append(f"{dt.get_text(' ', strip=True)}: {dd.get_text(' ', strip=True)}")
            if pairs:
                units.append(TextUnit("\n".join(pairs), {}, "html-dl", order=order))
                order += 1
            dl.decompose()
    # Whatever is left is the visible text. Joining elements with "\n" keeps neighbouring elements from running
    # together ("<b>Name</b><span>Priya</span>" must not become "NamePriya"); runs of blank lines are then
    # collapsed.
    text = soup.get_text("\n")
    text = re.sub(r"\n[ \t]*\n+", "\n", text)
    if text.strip():
        units.append(TextUnit(text, {}, "xml-text" if xml else "html-text", use_ner=True, order=order))
    return units


# ------------------------------------------------------------------ SQL dumps

_INSERT = re.compile(r"INSERT\s+(?:IGNORE\s+)?INTO\s+([`\"\[]?[\w.]+[`\"\]]?)\s*(?:\(([^)]*)\))?\s*VALUES\s*",
                     re.I)
_CREATE = re.compile(r"CREATE\s+TABLE\s+(?:IF\s+NOT\s+EXISTS\s+)?([`\"\[]?[\w.]+[`\"\]]?)\s*\((.*?)\)\s*;", re.I | re.S)
_COPY = re.compile(r"^COPY\s+([\w.\"]+)\s*\(([^)]*)\)\s+FROM\s+stdin;\s*$", re.I | re.M)


def _ident(s):
    return s.strip().strip("`\"[]").split(".")[-1]


def _sql_tuples(text, pos):
    """Parse '(v, 'x', NULL), (...);' starting at pos. Returns (list_of_rows, end_pos).

    A small hand-written scanner instead of a regex: quoted strings may contain commas, parentheses, doubled quotes
    ('O''Brien') and backslash escapes (MySQL dumps), all of which defeat a simple split."""
    rows, i, n = [], pos, len(text)
    while i < n:
        while i < n and text[i] in " \t\r\n,":
            i += 1
        if i >= n or text[i] != "(":
            break
        i += 1
        row, cur, inq = [], [], False
        while i < n:
            ch = text[i]
            if inq:
                if ch == "\\" and i + 1 < n:
                    cur.append(text[i + 1])
                    i += 2
                    continue
                if ch == "'":
                    if i + 1 < n and text[i + 1] == "'":
                        cur.append("'")
                        i += 2
                        continue
                    inq = False
                    i += 1
                    continue
                cur.append(ch)
            elif ch == "'":
                inq = True
            elif ch == ",":
                row.append("".join(cur).strip())
                cur = []
            elif ch == ")":
                row.append("".join(cur).strip())
                i += 1
                break
            else:
                cur.append(ch)
            i += 1
        rows.append([None if v.upper() == "NULL" else v for v in row])
        while i < n and text[i] in " \t\r\n":
            i += 1
        if i < n and text[i] == ";":
            return rows, i + 1
    return rows, i


def sql_units(text):
    """SQL dumps become tables: every INSERT / COPY row is a set of cells keyed by their column names, which is
    what tells a phone number from a timestamp in `VALUES (1, '9876543210', 1727584919793)`."""
    # Column names from CREATE TABLE, for INSERTs written without a column list. The split ignores commas inside
    # parentheses (numeric(10,2)); constraint lines are not columns.
    schemas = {}
    for m in _CREATE.finditer(text):
        cols = []
        for line in re.split(r",\s*(?![^()]*\))", m.group(2)):
            w = line.strip().split()
            if w and not re.match(r"(?i)(primary|unique|key|constraint|index|foreign|check)\b", w[0]):
                cols.append(_ident(w[0]))
        schemas[_ident(m.group(1))] = cols
    units, covered, order = [], [], 0
    for m in _INSERT.finditer(text):
        table = _ident(m.group(1))
        cols = [_ident(c) for c in m.group(2).split(",")] if m.group(2) else schemas.get(table, [])
        rows, end = _sql_tuples(text, m.end())
        covered.append((m.start(), end))
        units.extend(table_units([cols] + [[r or "" for r in row] for row in rows] if cols else
                                 [[r or "" for r in row] for row in rows], {}, "sql-insert", order=order))
        order += 1
    for m in _COPY.finditer(text):
        # PostgreSQL COPY ... FROM stdin: tab-separated rows up to a line holding "\.", with \N for NULL
        cols = [_ident(c) for c in m.group(2).split(",")]
        end = text.find("\n\\.", m.end())
        end = len(text) if end == -1 else end
        body = text[m.end():end]
        rows = [[None if v == "\\N" else v for v in line.split("\t")] for line in body.splitlines() if line]
        covered.append((m.start(), end))
        units.extend(table_units([cols] + [[r or "" for r in row] for row in rows], {}, "sql-copy", order=order))
        order += 1
    # Everything outside INSERT/COPY statements (comments, SET statements, GRANTs with passwords) is scanned as plain
    # text. The statements themselves are cut out so their values are not reported twice.
    covered.sort()
    rest, last = [], 0
    for s, e in covered:
        rest.append(text[last:s])
        last = max(last, e)
    rest.append(text[last:])
    other = "\n".join(r for r in rest if r.strip())
    if other.strip():
        units.append(TextUnit(other, {}, "sql-text", order=order))
    return units
