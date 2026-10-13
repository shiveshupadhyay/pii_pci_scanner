"""Build a labelled hidden-like test set (evaluator format) to estimate generalisation beyond dev.

    python scripts/make_synthetic.py --out synthetic [--seed 1] [--oov-share 0.5]
    python scan.py --input-dir synthetic/files --output-dir out/synth
    python starter/evaluate.py --labels synthetic/labels --output out/synth

Contents: formats absent from dev (TSV, HTML, SQL dumps, HTTP access logs, support tickets, .toml/.properties),
names never seen in dev (half of them outside the scanner's lexicons), hard negatives (Luhn-valid order ids,
test cards, invalid SSN/PAN/Aadhaar, hashes, UUIDs, placeholders, company/product phrases), misleading extensions,
OCR images, a column-only 60k-row table and unreadable files. Labels follow ANNOTATION_GUIDE.md.
Caveat: generator and scanner share an author, so this measures robustness, not annotation agreement.
"""

import argparse
import csv
import json
import os
import shutil
from email.message import EmailMessage
from email.policy import SMTP

from synth_lib import COMPANIES, PHRASES, PRODUCTS, Gold, Values, render_lines, save_jpeg, save_png, scanned_pdf


class Builder:
    def __init__(self, out, seed, oov_share, pool=None):
        self.files = os.path.join(out, "files")
        self.labels = os.path.join(out, "labels")
        self.v = Values(seed, oov_share, pool)
        self.r = self.v.r
        self.g = Gold()

    def path(self, rel, kind):
        self.g.paths[rel] = kind
        p = os.path.join(self.files, rel)
        os.makedirs(os.path.dirname(p), exist_ok=True)
        return p

    # ======================================================================= structured
    def table_spec(self, n, uninformative=False):
        """Columns of a generated table: (header, value generator, entity or None for a hard negative).

        With uninformative=True several headers become meaningless (col_7, ref_c, x9), which forces the scanner to
        recognise values by themselves. A Luhn-valid card under "col_7" must still be found, while the Luhn-valid
        order ids under "order_id" must not."""
        v, r = self.v, self.r
        card_col = "col_7" if uninformative else "card_number"
        cards = [v.card() for _ in range(n)]
        spec = [
            ("full_name" if not uninformative else "attr_2", lambda i: v.name(), "PERSON_NAME" if not uninformative else None),
            ("email" if not uninformative else "val_3", lambda i: v.email(), "EMAIL"),
            ("mobile", lambda i: v.phone(formatted=False), "PHONE"),
            ("pan_no" if not uninformative else "ref_c", lambda i: v.pan(), "NATIONAL_ID"),
            ("ssn" if not uninformative else "x9", lambda i: v.ssn(), "NATIONAL_ID"),
            (card_col, lambda i: cards[i], "CARD_NUMBER"),
            ("exp_date", lambda i: v.expiry(), "CARD_EXPIRY"),
            ("cvv", lambda i: v.cvv(amex=len(cards[i].replace(" ", "").replace("-", "")) == 15), "CVV"),
            ("db_password", lambda i: v.password(), "PASSWORD"),
            ("api_token", lambda i: v.token(r.choice(["hex", "b64", "stripe_live", "github"])), "API_TOKEN"),
            ("aws_key_id" if not uninformative else "field_5", lambda i: v.aws(), "AWS_ACCESS_KEY"),
            ("order_id", lambda i: v.order_id(), None),
            ("trace_id", lambda i: v.uuid(), None),
            ("created_ms", lambda i: v.epoch_ms(), None),
            ("bytes_sent", lambda i: v.bytes10(), None),
            ("card_last4", lambda i: str(r.randint(1000, 9999)), None),
            ("file_sha256", lambda i: v.sha256(), None),
            ("support_email", lambda i: r.choice(["support@example.com", "help@test.com", "admin@domain.com"]), None),
            ("amount", lambda i: f"{r.uniform(5, 9000):.2f}", None),
        ]
        if uninformative:  # an uninformative name column carries no context: values must speak for themselves
            spec = [s for s in spec if s[0] != "attr_2"]
        r.shuffle(spec)
        return spec

    def fill(self, spec, n, rel, sheet="", column_only=False, blank_share=0.1):
        cols = [c for c, _, _ in spec]
        rows = [[""] * len(cols) for _ in range(n)]
        for j, (c, gen, ent) in enumerate(spec):
            real = 0
            for i in range(n):
                x = self.r.random()
                # ~10% of sensitive cells are blank, or hold a placeholder in secret columns ("REDACTED",
                # "${DB_PASSWORD}"); those are not labelled and must not be reported
                if ent and x < blank_share:
                    rows[i][j] = self.r.choice(["", self.v.placeholder() if c in ("db_password", "api_token") else ""])
                    continue
                val = gen(i)
                rows[i][j] = val
                if ent:
                    real += 1
                    if not column_only:
                        loc = {"sheet": sheet, "column": c} if sheet else {"column": c}
                        self.g.add(rel, ent, val, loc)
            if ent and real:
                self.g.col(rel, c, ent, sheet)
        return cols, rows

    def csv_like(self, rel, delimiter=",", n=30, uninformative=False, column_only=False):
        p = self.path(rel, "structured")
        spec = self.table_spec(n, uninformative)
        cols, rows = self.fill(spec, n, rel, column_only=column_only)
        with open(p, "w", encoding="utf-8", newline="") as f:
            w = csv.writer(f, delimiter=delimiter)
            w.writerow(cols)
            w.writerows(rows)
        if column_only:
            self.g.column_only.append(rel)

    def lookalike_csv(self, rel):
        """Only hard negatives: no findings, no sensitive columns."""
        p = self.path(rel, "structured")
        v, r = self.v, self.r
        with open(p, "w", encoding="utf-8", newline="") as f:
            w = csv.writer(f)
            w.writerow(["order_id", "invoice_no", "tracking_number", "request_id", "card_last4", "masked_card",
                        "support_email", "created_ms", "status", "ssn_masked"])
            for _ in range(25):
                w.writerow([v.order_id(), v.order_id(), v.order_id(), v.uuid(), r.randint(1000, 9999),
                            v.masked_card(), "billing@example.org", v.epoch_ms(), r.choice(["open", "closed"]),
                            f"XXX-XX-{r.randint(1000, 9999)}"])

    def xlsx(self, rel):
        import openpyxl

        p = self.path(rel, "structured")
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for sheet, unin in (("Clients", False), ("Ledger", True)):
            ws = wb.create_sheet(sheet)
            spec = self.table_spec(18, unin)
            keep = spec[:9]
            if any(c in ("exp_date", "cvv") for c, _, _ in keep):  # expiry/CVV columns need their card column
                keep += [s for s in spec[9:] if s[0] in ("card_number", "col_7")]
            spec = keep
            cols, rows = self.fill(spec, 18, rel, sheet=sheet)
            ws.append(cols)
            for row in rows:
                ws.append(row)
        wb.save(p)

    def parquet(self, rel, n=40, column_only=False):
        import pyarrow as pa
        import pyarrow.parquet as pq

        p = self.path(rel, "structured")
        spec = self.table_spec(n)
        cols, rows = self.fill(spec, n, rel, column_only=column_only)
        table = pa.table({c: [row[j] for row in rows] for j, c in enumerate(cols)})
        pq.write_table(table, p, row_group_size=20000)
        if column_only:
            self.g.column_only.append(rel)

    def json_records(self, rel, jsonl=False, n=15):
        p = self.path(rel, "structured")
        v, r, g = self.v, self.r, self.g
        recs = []
        seen_cols = set()

        def put(path, ent, val):
            g.add(rel, ent, val, {"column": path})
            seen_cols.add((path, ent))
            return val

        for i in range(n):
            name = v.name()
            cards = []
            for _ in range(r.randint(1, 3)):
                if r.random() < 0.2:
                    cards.append({"number": v.masked_card()})
                else:
                    c = v.card()
                    cards.append({"number": put("billing.cards[].number", "CARD_NUMBER", c),
                                  "expiry": put("billing.cards[].expiry", "CARD_EXPIRY", v.expiry())})
            rec = {"id": i + 1,
                   "profile": {"name": put("profile.name", "PERSON_NAME", name),
                               "contact": {"email": put("profile.contact.email", "EMAIL", v.email(name)),
                                           "phone": put("profile.contact.phone", "PHONE", v.phone())}},
                   "identity": {"aadhaar": put("identity.aadhaar", "NATIONAL_ID", v.aadhaar()),
                                "pan": put("identity.pan", "NATIONAL_ID", v.pan())},
                   "billing": {"cards": cards, "invoice_no": v.order_id()},
                   "auth": {"access_token": put("auth.access_token", "API_TOKEN", v.token("jwt")),
                            "token_sha": v.sha256()},
                   "meta": {"trace_id": v.uuid(), "created_ms": int(v.epoch_ms())}}
            recs.append(rec)
        for path, ent in sorted(seen_cols):
            g.col(rel, path, ent)
        with open(p, "w", encoding="utf-8", newline="\n") as f:
            if jsonl:
                for rec in recs:
                    f.write(json.dumps(rec) + "\n")
            else:
                json.dump(recs, f, indent=2)

    # ======================================================================= text documents
    def html(self, rel):
        p = self.path(rel, "text")
        v, g = self.v, self.g
        rows = []
        for _ in range(6):
            name, card = v.name(), v.card()
            row = [name, v.email(name), v.phone(), v.pan(), card, v.expiry()]
            for ent, val in zip(["PERSON_NAME", "EMAIL", "PHONE", "NATIONAL_ID", "CARD_NUMBER", "CARD_EXPIRY"], row):
                g.add(rel, ent, val)
            rows.append(row)
        neg = [v.order_id(), v.uuid(), COMPANIES[0]]
        holder, mob, pw = v.name(), v.phone(False), v.password()
        for ent, val in (("PERSON_NAME", holder), ("PHONE", mob)):
            g.add(rel, ent, val)
        body_name = v.name()
        g.add(rel, "PERSON_NAME", body_name)
        em = v.email()
        g.add(rel, "EMAIL", em)
        html = f"""<!DOCTYPE html><html><head><title>Customer export</title>
<style>td {{ padding: 4px }}</style></head><body>
<h1>Customer register &amp; payments</h1>
<p>Escalated by {body_name} &mdash; contact <a href="mailto:{em}">{em}</a>. Order {neg[0]}, trace {neg[1]}.</p>
<table><tr><th>Name</th><th>Email</th><th>Phone</th><th>PAN</th><th>Card Number</th><th>Expiry</th></tr>
{''.join('<tr>' + ''.join(f'<td>{c}</td>' for c in r) + '</tr>' for r in rows)}
</table>
<form><label for="h">Account holder</label><input id="h" name="holder" value="{holder}">
<label for="m">Mobile</label><input id="m" name="mobile" value="{mob}">
<input type="password" name="password" placeholder="Password"></form>
<p>Vendor: {neg[2]} &amp; {PRODUCTS[0]}. Questions? support@example.com</p>
</body></html>"""
        with open(p, "w", encoding="utf-8", newline="\n") as f:
            f.write(html)

    def sql_dump(self, rel):
        p = self.path(rel, "text")
        v, g = self.v, self.g
        lines = ["-- PostgreSQL database dump", "SET client_encoding = 'UTF8';",
                 "CREATE TABLE public.customers (id integer, full_name text, email text, phone text, "
                 "aadhaar_no text, card_number text, password_hash text, created_ms bigint);"]
        vals = []
        for i in range(8):
            name = v.name()
            row = [str(i + 1), name, v.email(name), v.phone(False), v.aadhaar(spaced=False), v.card("visa", "plain"),
                   v.bcrypt(), v.epoch_ms()]
            for ent, val in zip([None, "PERSON_NAME", "EMAIL", "PHONE", "NATIONAL_ID", "CARD_NUMBER", None, None], row):
                if ent:
                    g.add(rel, ent, val)
            vals.append("(" + ", ".join(row[:1] + [f"'{x}'" for x in row[1:7]] + row[7:]) + ")")
        lines.append("INSERT INTO public.customers (id, full_name, email, phone, aadhaar_no, card_number, "
                     "password_hash, created_ms) VALUES\n  " + ",\n  ".join(vals) + ";")
        lines.append("CREATE TABLE app_users (uid int, login text, api_token text, pwd text);")
        rows2 = []
        for i in range(5):
            tok, pw = v.token("hex"), v.password()
            g.add(rel, "API_TOKEN", tok)
            g.add(rel, "PASSWORD", pw)
            rows2.append(f"({i}, 'user{i}', '{tok}', '{pw}')")
        lines.append("INSERT INTO app_users VALUES " + ", ".join(rows2) + ";")
        lines.append("COPY public.orders (order_id, tracking_number, buyer_email) FROM stdin;")
        for _ in range(5):
            em = v.email()
            g.add(rel, "EMAIL", em)
            lines.append(f"{v.order_id()}\t{v.order_id()}\t{em}")
        lines.append("\\.")
        with open(p, "w", encoding="utf-8", newline="\n") as f:
            f.write("\n".join(lines) + "\n")

    def http_log(self, rel, n=400):
        p = self.path(rel, "text")
        v, r, g = self.v, self.r, self.g
        out = []
        for i in range(n):
            ip = f"{r.randint(11, 220)}.{r.randint(0, 255)}.{r.randint(0, 255)}.{r.randint(1, 254)}"
            ts = f"[{r.randint(10, 28)}/Oct/2026:{r.randint(10, 23)}:{r.randint(10, 59)}:{r.randint(10, 59)} +0000]"
            k = r.random()
            if k < 0.05:
                em = v.email()
                g.add(rel, "EMAIL", em)
                path = f"/api/v1/users?email={em}&page={r.randint(1, 9)}"
            elif k < 0.09:
                tok = v.token(r.choice(["stripe_live", "github", "google"]))
                g.add(rel, "API_TOKEN", tok)
                path = f"/v1/charges?api_key={tok}&limit=10"
            elif k < 0.11:
                tok = v.token("hex")
                g.add(rel, "API_TOKEN", tok)
                path = f"/export?token={tok}"
            else:
                path = r.choice(["/static/app.3f9a1c2b.js", f"/api/v1/orders/{r.randint(1000, 99999)}", "/health",
                                 f"/api/v1/items/{v.order_id()}"])
            out.append(f'{ip} - - {ts} "GET {path} HTTP/1.1" {r.choice([200, 200, 304, 404, 500])} '
                       f'{r.randint(200, 90000)} "-" "Mozilla/5.0 (X11; Linux x86_64) Chrome/{r.randint(90, 130)}.0"')
            if r.random() < 0.04:
                jwt = v.token("jwt")
                g.add(rel, "API_TOKEN", jwt)
                out.append(f"    Authorization: Bearer {jwt}")
            if r.random() < 0.03:
                sess = v.token("b64")
                g.add(rel, "API_TOKEN", sess)
                out.append(f"    Cookie: session={sess}; theme=dark; lang=en")
            if r.random() < 0.03:
                out.append(f"    X-Request-Id: {v.uuid()}")
        with open(p, "w", encoding="utf-8", newline="\n") as f:
            f.write("\n".join(out) + "\n")

    def ticket(self, rel):
        p = self.path(rel, "text")
        v, r, g = self.v, self.r, self.g
        out = []
        for t in range(3):
            req, agent = v.name(), v.name()
            em, ph = v.email(req), v.phone(False)
            card, exp = v.card(fmt="space"), v.expiry()
            for ent, val in (("PERSON_NAME", req), ("EMAIL", em), ("PHONE", ph), ("CARD_NUMBER", card),
                             ("CARD_EXPIRY", exp), ("PERSON_NAME", agent)):
                g.add(rel, ent, val)
            pan = v.pan()
            g.add(rel, "NATIONAL_ID", pan)
            out += [f"Ticket #{r.randint(10000, 99999)} | Priority: {r.choice(['High', 'Low'])} | Status: Open",
                    f"Requester: {req} <{em}>",
                    f"Phone: {ph}",
                    f"Created: 2026-0{r.randint(1, 9)}-1{r.randint(0, 9)} 10:2{r.randint(0, 9)}",
                    "Description:",
                    f"Hi team, my card {card} (expiry {exp}) was charged twice for order {v.order_id()}.",
                    f"Our {r.choice(PHRASES)} contact at {r.choice(COMPANIES)} said to raise a ticket.",
                    "--",
                    f"Assignee: {agent}",
                    f"Internal note: verified PAN {pan}; legacy id {v.bad_pan()}; ref {v.bad_aadhaar()}.",
                    ""]
        with open(p, "w", encoding="utf-8", newline="\n") as f:
            f.write("\n".join(out))

    def app_log(self, rel, n=300):
        p = self.path(rel, "text")
        v, r, g = self.v, self.r, self.g
        out = []
        for i in range(n):
            ts = f"2026-0{r.randint(1, 9)}-{r.randint(10, 28)}T{r.randint(10, 23)}:{r.randint(10, 59)}:{r.randint(10, 59)}Z"
            lvl = r.choice(["INFO", "WARN", "DEBUG", "ERROR"])
            k = r.random()
            if k < 0.04:
                nm = v.name()
                g.add(rel, "PERSON_NAME", nm)
                msg = f"kyc approved for {nm} (customer={r.randint(1000, 9999)})"
            elif k < 0.08:
                c, e, cv = v.card(), v.expiry(), None
                g.add(rel, "CARD_NUMBER", c)
                g.add(rel, "CARD_EXPIRY", e)
                amex = len(c.replace(" ", "").replace("-", "")) == 15
                cv = v.cvv(amex)
                g.add(rel, "CVV", cv)
                msg = f"charge failed pan={c} exp={e} cvc={cv}"
            elif k < 0.11:
                ssn = v.ssn()
                g.add(rel, "NATIONAL_ID", ssn)
                msg = f"tax lookup ssn={ssn} ok"
            elif k < 0.13:
                pw = v.password()
                g.add(rel, "PASSWORD", pw)
                msg = f"db connect redis://default:{pw}@cache-3.svc:6379/0"
            elif k < 0.15:
                ph = v.phone()
                g.add(rel, "PHONE", ph)
                msg = f"whatsapp notification queued to {ph}"
            elif k < 0.17:
                ak = v.aws()
                g.add(rel, "AWS_ACCESS_KEY", ak)
                msg = f"assumed role with key {ak}"
            else:
                msg = r.choice([f"order {v.order_id()} shipped", f"ssn check skipped for {v.bad_ssn()}",
                                f"retry {r.randint(1, 5)}/5 status=503", f"build {r.randint(1, 9)}.{r.randint(0, 20)}.{r.randint(0, 99)}",
                                f"cache hit key=user:{r.randint(1, 99999)}", f"latency_ms={r.randint(1, 900)}",
                                f"card {r.choice(['4111111111111111', '4242424242424242'])} used in sandbox",
                                f"checksum={v.md5()}", f"request_id={v.uuid()}", f"bytes_out={v.bytes10()}",
                                f"user {r.choice(PHRASES)} updated", f"masked card {v.masked_card()}",
                                f"AWS_ACCESS_KEY_ID=AKIAIOSFODNN7EXAMPLE", f"password_min_length = {r.randint(8, 14)}"])
            out.append(f"{ts} {lvl} [svc-{r.randint(1, 5)}] {msg}")
        with open(p, "w", encoding="utf-8", newline="\n") as f:
            f.write("\n".join(out) + "\n")

    def configs(self):
        v, g = self.v, self.g
        rel = "conf/service.toml"
        p = self.path(rel, "text")
        pw, tok, ak = v.password(), v.token("stripe_live"), v.aws()
        g.add(rel, "PASSWORD", pw)
        g.add(rel, "API_TOKEN", tok)
        g.add(rel, "AWS_ACCESS_KEY", ak)
        with open(p, "w", encoding="utf-8", newline="\n") as f:
            f.write(f'[database]\nhost = "db.internal"\npassword = "{pw}"\npassword_rotation_days = 90\n\n'
                    f'[payments]\nstripe_secret = "{tok}"\nwebhook_url = "https://hooks.example.com/x"\n\n'
                    f'[aws]\naccess_key_id = "{ak}"\nregion = "eu-west-1"\nsecret_ref = "${{AWS_SECRET}}"\n')
        rel = "conf/app.properties"
        p = self.path(rel, "text")
        pw2, tok2 = v.password(), v.token("b64")
        g.add(rel, "PASSWORD", pw2)
        g.add(rel, "API_TOKEN", tok2)
        with open(p, "w", encoding="utf-8", newline="\n") as f:
            f.write(f"spring.datasource.password={pw2}\nspring.datasource.username=app\n"
                    f"client.secret={tok2}\nbuild.commit={v.sha256()[:40]}\nadmin.password=${{ADMIN_PASSWORD}}\n")
        rel = "repo/worker.go"
        p = self.path(rel, "text")
        tok3 = v.token("github")
        author = v.name()
        em = v.email(author)
        g.add(rel, "API_TOKEN", tok3)
        g.add(rel, "PERSON_NAME", author)
        g.add(rel, "EMAIL", em)
        with open(p, "w", encoding="utf-8", newline="\n") as f:
            f.write(f'// Maintainer: {author} <{em}>\npackage main\n\nconst ghToken = "{tok3}"\n'
                    f'var apiKey = os.Getenv("API_KEY")\nvar testCard = "4242424242424242"\n'
                    f'// sha256: {v.sha256()}\n')

    def pem_files(self):
        rel = "keys/deploy.key"
        p = self.path(rel, "text")
        k = self.v.private_key()
        self.g.add(rel, "PRIVATE_KEY", k)
        with open(p, "w", encoding="utf-8", newline="\n") as f:
            f.write(k + "\n")
        rel = "keys/ca.crt"
        p = self.path(rel, "text")
        body = "\n".join("".join(self.r.choice("ABCDEFabcdef0123456789+/") for _ in range(64)) for _ in range(8))
        with open(p, "w", encoding="utf-8", newline="\n") as f:
            f.write(f"-----BEGIN CERTIFICATE-----\n{body}\n-----END CERTIFICATE-----\n")

    def emails(self):
        v, g = self.v, self.g
        for i, cte in enumerate(["base64", "quoted-printable", "alternative"]):
            rel = f"mail/msg_{i}.eml"
            p = self.path(rel, "text")
            a, b = v.name(), v.name()
            ea, eb = v.email(a), v.email(b)
            for ent, val in (("PERSON_NAME", a), ("EMAIL", ea), ("PERSON_NAME", b), ("EMAIL", eb)):
                g.add(rel, ent, val)
            card, exp, pw = v.card(fmt="dash"), v.expiry(), v.password()
            third = v.name()
            body = (f"Hi {b.split()[0]},\n\nYour temporary password is {pw} - please change it after login.\n"
                    f"We charged card {card}, valid thru {exp}. I've cc'd {third} from {COMPANIES[1]}.\n\n"
                    f"Thanks,\n{a.split()[0]}\n{PHRASES[0]} | {COMPANIES[1]}\n")
            for ent, val in (("PASSWORD", pw), ("CARD_NUMBER", card), ("CARD_EXPIRY", exp), ("PERSON_NAME", third)):
                g.add(rel, ent, val)
            m = EmailMessage(policy=SMTP)
            m["From"], m["To"], m["Subject"] = f"{a} <{ea}>", f"{b} <{eb}>", "Account update"
            # a Message-ID shaped like an address on the sender's real domain: the classic e-mail false positive
            m["Message-ID"] = f"<{v.md5()[:16]}@mail.{ea.split('@')[1]}>"
            if cte == "alternative":
                m.set_content(body)
                m.add_alternative("<html><body><p>" + body.replace("\n", "<br>") + "</p></body></html>", subtype="html")
                # the HTML part repeats every body value: each occurrence counts
                for ent, val in (("PASSWORD", pw), ("CARD_NUMBER", card), ("CARD_EXPIRY", exp), ("PERSON_NAME", third)):
                    g.add(rel, ent, val)
            else:
                m.set_content(body, cte=cte)
            with open(p, "wb") as f:
                f.write(m.as_bytes())

    def docx(self, rel):
        import docx

        p = self.path(rel, "text")
        v, g = self.v, self.g
        d = docx.Document()
        d.add_heading("Vendor onboarding", 1)
        a = v.name()
        g.add(rel, "PERSON_NAME", a)
        d.add_paragraph(f"Prepared by: {a}")
        d.add_paragraph(f"Approved for {COMPANIES[2]} under {PRODUCTS[1]} licence. Reference {v.order_id()}.")
        t = d.add_table(rows=1, cols=5)
        for c, h in zip(t.rows[0].cells, ["Contact", "Email", "Mobile", "Aadhaar", "IBAN ref"]):
            c.text = h
        for _ in range(5):
            n = v.name()
            row = [n, v.email(n), v.phone(False), v.aadhaar(), v.bad_aadhaar()]
            cells = t.add_row().cells
            for c, val in zip(cells, row):
                c.text = val
            for ent, val in zip(["PERSON_NAME", "EMAIL", "PHONE", "NATIONAL_ID"], row):
                g.add(rel, ent, val)
        d.save(p)

    def pdf(self, rel, as_name=None):
        import pymupdf

        p = self.path(as_name or rel, "text")
        rel = as_name or rel
        v, r, g = self.v, self.r, self.g
        doc = pymupdf.open()
        # page 1: form
        pg = doc.new_page(width=595, height=842)
        pg.insert_text((50, 50), "Applicant details", fontsize=12)
        name, mob, card, exp = v.name(), v.phone(False), v.card(fmt="space"), v.expiry()
        fields = [("Full name", name, "PERSON_NAME"), ("Mobile", mob, "PHONE"), ("SSN", v.ssn(), "NATIONAL_ID"),
                  ("Card number", card, "CARD_NUMBER"), ("Valid thru", exp, "CARD_EXPIRY"),
                  ("Order ID", v.order_id(), None), ("Reference", v.uuid(), None)]
        for i, (lab, val, ent) in enumerate(fields):
            y = 85 + i * 18
            pg.insert_text((50, y), lab, fontsize=9)
            pg.insert_text((150, y), val, fontsize=9)
            if ent:
                g.add(rel, ent, val, {"page": 1})
        # page 2: table
        pg = doc.new_page(width=595, height=842)
        pg.insert_text((45, 45), "Payout register", fontsize=12)
        hdr = [("payee", 45), ("phone", 130), ("bank_ref", 215), ("card", 300), ("expiry", 400), ("txn_id", 460)]
        for h, x in hdr:
            pg.insert_text((x, 75), h, fontsize=8)
        for i in range(10):
            y = 88 + i * 12
            n, ph, c, e = v.name(), v.phone(False), v.card(fmt="plain"), v.expiry()
            row = [n, ph, v.bad_aadhaar(), c, e, v.order_id()]
            for (h, x), val in zip(hdr, row):
                pg.insert_text((x, y), val, fontsize=8)
            for ent, val in (("PERSON_NAME", n), ("PHONE", ph), ("CARD_NUMBER", c), ("CARD_EXPIRY", e)):
                g.add(rel, ent, val, {"page": 2})
        # page 3: two columns of prose
        pg = doc.new_page(width=595, height=842)
        left_name, right_name = v.name(), v.name()
        lem, rpw = v.email(left_name), v.password()
        left = f"Escalation owner {left_name} can be reached at {lem}. The {PHRASES[3]} is next week."
        right = f"Temporary credentials were shared: the new password is {rpw} - rotate after use."
        pg.insert_textbox(pymupdf.Rect(50, 60, 280, 300), left, fontsize=9)
        pg.insert_textbox(pymupdf.Rect(320, 60, 550, 300), right, fontsize=9)
        for ent, val in (("PERSON_NAME", left_name), ("EMAIL", lem), ("PASSWORD", rpw)):
            g.add(rel, ent, val, {"page": 3})
        doc.save(p)
        doc.close()

    # ======================================================================= OCR
    def screenshot(self, rel, dark=False, font="mono", fmt="png", kind_of="config"):
        p = self.path(rel, "ocr")
        v, g = self.v, self.g
        loc = {"page": 1}
        if kind_of == "config":
            pw, tok, em = v.password(), v.token(self.r.choice(["stripe_test", "gitlab", "hex"])), v.email()
            pw = "".join(ch for ch in pw if ch not in "Il1O0|")  # keep the test about pipeline, not glyph ambiguity
            tok = tok[:8] + "".join(ch for ch in tok[8:] if ch not in "Il1O0|")
            lines = [f"DB_PASSWORD={pw}", f"API_TOKEN={tok}", f"ADMIN_EMAIL={em}", "REPLICAS=3",
                     f"BUILD_SHA={v.md5()}", "CACHE_TTL=300"]
            for ent, val in (("PASSWORD", pw), ("API_TOKEN", tok), ("EMAIL", em)):
                g.add(rel, ent, val, loc)
        else:
            a, b = v.name(), v.name()
            ea, eb = v.email(a), v.email(b)
            ph = v.phone()
            lines = [f"From: {a} <{ea}>", f"To: {b} <{eb}>", "Subject: Quick question", "",
                     f"Hi {b.split()[0]},", f"Please call me back on {ph} about order {v.order_id()}.", "",
                     "Thanks,", a.split()[0], f"{PHRASES[1]} | {COMPANIES[3]}"]
            for ent, val in (("PERSON_NAME", a), ("EMAIL", ea), ("PERSON_NAME", b), ("EMAIL", eb), ("PHONE", ph)):
                g.add(rel, ent, val, loc)
        img = render_lines(lines, font=font, size=22, dark=dark, noise=0.0 if fmt == "png" else 0.004,
                           seed=self.r.randint(0, 10 ** 6))
        if fmt == "png":
            save_png(img, p)
        else:
            save_jpeg(img, p)

    def scanned(self, rel):
        p = self.path(rel, "ocr")
        v, g = self.v, self.g
        imgs = []
        a = v.name()
        ea = v.email(a)
        card = v.card(fmt="space")
        lines1 = [f"Statement for {a}", f"Email: {ea}", f"Card: {card}", f"Order ID: {v.order_id()}",
                  "Status: settled"]
        for ent, val in (("PERSON_NAME", a), ("EMAIL", ea), ("CARD_NUMBER", card)):
            g.add(rel, ent, val, {"page": 1})
        imgs.append(render_lines(lines1, font="serif", size=24, noise=0.003, seed=1))
        lines2 = ["# Release Notes", f"# version {self.r.randint(1, 3)}.{self.r.randint(0, 9)}.{self.r.randint(0, 99)}",
                  "# paid with card ending in 4242", f"# checksum: {v.md5()}", f"# SSN {v.bad_ssn()}"]
        imgs.append(render_lines(lines2, font="sans", size=24, noise=0.003, seed=2))
        scanned_pdf(imgs, p)

    # ======================================================================= errors
    def errors(self):
        """One file of each unreadable kind: empty, truncated PDF, encrypted PDF, corrupt XLSX (a ZIP header
        followed by garbage), and random bytes with a .docx name."""
        import pymupdf

        g = self.g
        for rel in ("exports/empty_export.csv",):
            open(self.path(rel, "error"), "wb").close()
            g.errors.append(rel)
        rel = "contracts/broken.pdf"
        doc = pymupdf.open()
        doc.new_page().insert_text((50, 50), "draft")
        data = doc.tobytes()
        with open(self.path(rel, "error"), "wb") as f:
            f.write(data[: len(data) // 3])
        g.errors.append(rel)
        rel = "finance/locked.pdf"
        doc = pymupdf.open()
        doc.new_page().insert_text((50, 50), f"salary of {self.v.name()}")
        doc.save(self.path(rel, "error"), encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="u53r", owner_pw="0wn3r")
        g.errors.append(rel)
        rel = "finance/budget.xlsx"
        with open(self.path(rel, "error"), "wb") as f:
            f.write(b"PK\x03\x04" + os.urandom(3000))
        g.errors.append(rel)
        rel = "hr/minutes.docx"
        with open(self.path(rel, "error"), "wb") as f:
            f.write(os.urandom(5000))
        g.errors.append(rel)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default="synthetic")
    ap.add_argument("--seed", type=int, default=1)
    ap.add_argument("--oov-share", type=float, default=0.5)
    ap.add_argument("--pool", choices=["train", "test"], default="test",
                    help="name-pool half: the name model is trained on 'train', evaluated on 'test'")
    a = ap.parse_args()
    if os.path.exists(a.out):
        shutil.rmtree(a.out)
    b = Builder(a.out, a.seed, a.oov_share, a.pool)
    b.csv_like("exports/customers.csv")
    b.csv_like("exports/uninformative.csv", uninformative=True)
    b.csv_like("exports/partners.tsv", delimiter="\t")
    b.csv_like("exports/partners_tab.txt", delimiter="\t", n=20)      # TSV with a misleading extension
    b.csv_like("warehouse/events_big.csv", n=60_000, column_only=True)
    b.lookalike_csv("exports/order_refs.csv")
    b.xlsx("hr/staff.xlsx")
    b.parquet("lake/accounts.parquet")
    b.parquet("lake/events_big.parquet", n=60_000, column_only=True)
    b.json_records("api/customers.json")
    b.json_records("api/customers.jsonl", jsonl=True)
    b.json_records("api/stream.log", jsonl=True)                     # JSONL named .log
    b.html("web/customers.html")
    b.html("web/report.htm")
    b.sql_dump("db/dump.sql")
    b.http_log("logs/access.log")
    b.ticket("support/ticket_1001.txt")
    b.ticket("support/ticket_1002.txt")
    b.app_log("logs/app.log")
    b.configs()
    b.pem_files()
    b.emails()
    b.docx("shared/onboarding.docx")
    b.pdf("reports/payouts.pdf")
    b.pdf(None, as_name="reports/payouts_copy.bin")                  # PDF with a misleading extension
    b.screenshot("screens/config_light.png", dark=False, font="mono")
    b.screenshot("screens/config_dark.png", dark=True, font="dejavumono")
    b.screenshot("screens/mail.png", font="sans", kind_of="mail")
    b.screenshot("scans/mail_scan.jpg", font="serif", fmt="jpg", kind_of="mail")
    b.screenshot("screens/mail_shot.pdf", font="noto", kind_of="mail")  # PNG named .pdf
    b.scanned("scans/statement_scan.pdf")
    b.errors()
    b.g.write(b.labels)
    print(f"wrote {len(b.g.paths)} files, {len(b.g.findings)} findings, {len(b.g.columns)} columns -> {a.out}")


if __name__ == "__main__":
    main()
