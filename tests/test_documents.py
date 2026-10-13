"""Document extraction keeps layout context: table headers, form labels, two-column pages, MIME decoding."""

from email.message import EmailMessage
from email.policy import SMTP

from scanner.extract import documents
from scanner.extract.pdf import pdf_pages


def detect_units(det, units):
    out = []
    for u in units:
        out += [(c.entity, c.value) for c in det.detect(u.text, key=u.key, cell=u.cell, table=u.table)]
    return out


def test_docx_table_header_context(tmp_path, det):
    import docx

    d = docx.Document()
    d.add_paragraph("From: Arjun Singh <arjun_singh@yahoo.co.in>")
    t = d.add_table(rows=1, cols=3)
    for c, h in zip(t.rows[0].cells, ["ssn", "phone_no", "bytes_sent"]):
        c.text = h
    cells = t.add_row().cells
    for c, v in zip(cells, ["743-25-1262", "8029080667", "8043983594"]):
        c.text = v
    p = tmp_path / "a.docx"
    d.save(p)
    found = detect_units(det, documents.docx_units(str(p)))
    assert ("PHONE", "8029080667") in found and ("NATIONAL_ID", "743-25-1262") in found
    assert ("PHONE", "8043983594") not in found
    assert ("PERSON_NAME", "Arjun Singh") in found


def test_eml_decoding_base64_and_soft_line_breaks(tmp_path, det):
    for cte in ("base64", "quoted-printable"):
        m = EmailMessage(policy=SMTP)
        m["From"], m["To"], m["Subject"] = "Ananya Iyer <ananya.iyer@northwind.io>", "Aarav Kulkarni <ak@outlook.com>", "x"
        m["Message-ID"] = "<bd687727e0d808d1@mail.outlook.com>"
        m.set_content("Hi Aarav,\n\n" + "filler " * 12 + "charged to card 6012-2118-9541-8701, valid thru 11/29.\n",
                      cte=cte)
        p = tmp_path / f"m_{cte}.eml"
        p.write_bytes(m.as_bytes())
        found = detect_units(det, documents.eml_units(str(p)))
        assert ("CARD_NUMBER", "6012-2118-9541-8701") in found and ("CARD_EXPIRY", "11/29") in found
        assert ("EMAIL", "ananya.iyer@northwind.io") in found and ("PERSON_NAME", "Aarav Kulkarni") in found
        assert not any(v.startswith("bd687727") for _, v in found)


def test_html_entities_tables_and_forms(det):
    html = """<html><body><p>Fish &amp; Chips Ltd &mdash; contact ravi.k&#64;orbitpay.in</p>
    <table><tr><th>Name</th><th>Mobile</th><th>Card Number</th><th>Expiry</th></tr>
    <tr><td>Meera Nair</td><td>9876501234</td><td>4539148803436467</td><td>08/27</td></tr></table>
    <label for="m">Mobile</label><input id="m" value="9123456780"></body></html>"""
    found = detect_units(det, documents.html_units(html))
    for v in [("EMAIL", "ravi.k@orbitpay.in"), ("PERSON_NAME", "Meera Nair"), ("PHONE", "9876501234"),
              ("CARD_NUMBER", "4539148803436467"), ("CARD_EXPIRY", "08/27"), ("PHONE", "9123456780")]:
        assert v in found


def test_sql_insert_column_context(det):
    sql = ("CREATE TABLE users (id int, full_name text, phone text, created_ms bigint);\n"
           "INSERT INTO users VALUES (1, 'Priya Sharma', '9876543210', 1727584919793), "
           "(2, 'O''Brien Kelly', '9123456780', 1727584919794);\n-- password = Adm1n405!\n")
    found = detect_units(det, documents.sql_units(sql))
    assert ("PHONE", "9876543210") in found and ("PERSON_NAME", "Priya Sharma") in found
    assert ("PASSWORD", "Adm1n405!") in found
    assert not any(v == "1727584919793" for _, v in found)


def _pdf_form(path, two_columns=False):
    # A label/value form as reports generate them: labels at x=50, values at x=150. With two_columns=True a second
    # form at x=310/410 mirrors it on the right half of the page, so the page has to be split at the gutter.
    import pymupdf

    doc = pymupdf.open()
    pg = doc.new_page(width=595, height=842)
    rows = [("Full name", "Priya Wang"), ("Mobile", "8424350180"), ("Card number", "5358 6347 6230 0815"),
            ("Valid thru", "06/30"), ("Order ID", "4322373367923988")]
    for i, (lab, val) in enumerate(rows):
        pg.insert_text((50, 70 + i * 18), lab, fontsize=9)
        pg.insert_text((150, 70 + i * 18), val, fontsize=9)
        if two_columns:
            pg.insert_text((310, 70 + i * 18), lab, fontsize=9)
            pg.insert_text((410, 70 + i * 18), "Emily Okafor" if lab == "Full name" else
                           ("9765389200" if lab == "Mobile" else val), fontsize=9)
    doc.save(path)
    doc.close()


def test_pdf_form_labels(tmp_path, det):
    p = str(tmp_path / "form.pdf")
    _pdf_form(p)
    found = []
    for kind, pno, units in pdf_pages(p):
        found += detect_units(det, units)
    assert ("PHONE", "8424350180") in found and ("CARD_EXPIRY", "06/30") in found
    assert ("PERSON_NAME", "Priya Wang") in found
    assert not any(v == "4322373367923988" for _, v in found)  # Order ID label


def test_pdf_two_columns_are_separate_contexts(tmp_path, det):
    p = str(tmp_path / "two.pdf")
    _pdf_form(p, two_columns=True)
    kinds, found = [], []
    for kind, pno, units in pdf_pages(p):
        kinds += [u.extractor for u in units]
        found += detect_units(det, units)
    assert ("PHONE", "9765389200") in found and ("PERSON_NAME", "Emily Okafor") in found
    assert sum(1 for k in kinds if k.startswith("pdf-")) >= 2


def test_pdf_table_header_context(tmp_path, det):
    import pymupdf

    doc = pymupdf.open()
    pg = doc.new_page(width=595, height=842)
    # Tight columns (status and phone only 35 pt apart at 7 pt type) and a header with digits (file_sha256): the two
    # things that broke the first table detector. Each value starts exactly at its header's x position.
    hdr = [("status", 45), ("phone", 80), ("file_sha256", 130), ("password", 420), ("employee", 490)]
    for h, x in hdr:
        pg.insert_text((x, 76), h, fontsize=7)
    vals = [["pending", "9586992672", "49620b50e0ea3f28a4b0f3ebb9850d4917395c6d558e6162835c11eb8af2d604", "aqEZgtuwlz4D",
             "Meera Khan"], ["active", "7152038462", "0a63194e13e59eeb6b3134e463bd09ab34dd7a6d935c1ee202bb2b8f9b2b5cb4",
                             "Welcome414", "Ethan Miller"]]
    for i, row in enumerate(vals):
        for (h, x), v in zip(hdr, row):
            pg.insert_text((x, 88 + i * 9), v, fontsize=7)
    p = str(tmp_path / "table.pdf")
    doc.save(p)
    found = []
    for kind, pno, units in pdf_pages(p):
        found += detect_units(det, units)
    for v in [("PHONE", "9586992672"), ("PASSWORD", "aqEZgtuwlz4D"), ("PERSON_NAME", "Meera Khan"),
              ("PASSWORD", "Welcome414")]:
        assert v in found
    assert not any(e == "API_TOKEN" for e, _ in found)  # sha256 column is a hash, not a token
